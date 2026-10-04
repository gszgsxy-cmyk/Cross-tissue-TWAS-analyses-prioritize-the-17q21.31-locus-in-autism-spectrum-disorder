import pathlib, re, json, hashlib, math
from collections import defaultdict
import pandas as pd
import openpyxl
def fisher_exact(table,alternative='greater'):
 a,b=table[0];c,d=table[1]; n=a+b+c+d; k=a+c; draws=a+b
 def lc(n,k):return math.lgamma(n+1)-math.lgamma(k+1)-math.lgamma(n-k+1)
 prob=sum(math.exp(lc(k,x)+lc(n-k,draws-x)-lc(n,draws)) for x in range(a,min(k,draws)+1))
 return a*d/(b*c) if b*c else float('inf'),min(1,prob)

ROOT=pathlib.Path('__WORKSPACE_ROOT__')
WORK=ROOT/'work/external_benchmark_20260926'
OUT=ROOT/'outputs/external_benchmark_20260926'; OUT.mkdir(exist_ok=True)
DL=pathlib.Path('__USER_HOME__/Downloads')
TW=pathlib.Path('__USER_HOME__/Desktop/TWAS')
paths={'JTI':TW/'MR-JTI-master/model_training/predixcan/src/ASD.JTI.29tissues.updated.all_withFDR.tsv','FUSION':TW/'FUSION/fusion_twas-master/ASD.FUSION.all29.GTExv8WB.updated.GENCODEv26.allchr.withFDR.dat','GENCODE':TW/'FUSION/fusion_twas-master/gencode.v26.GRCh38.genes.gtf'}
names={'HuangS1':'journal.pgen.1009309.s024.xlsx','HuangS2':'journal.pgen.1009309.s025.xlsx','HuangS3':'journal.pgen.1009309.s026.xlsx','Matoba':'41398_2020_953_MOESM5_ESM.xlsx','Gandal':'41586_2022_5377_MOESM5_ESM.xlsx'}
paths.update({k:DL/v for k,v in names.items()})
inventory=[]; unmatched=[]; sheets={}; symids=defaultdict(set); ann={}
for line in paths['GENCODE'].open(encoding='utf8'):
 if line.startswith('#'):continue
 c=line.rstrip().split('\t'); a=dict(re.findall(r'(\w+) "([^"]*)"',c[8])); e=a['gene_id'].split('.')[0]; s=a['gene_name']
 ann[e]={'symbol':s,'chr':c[0],'start':int(c[3]),'end':int(c[4])}; symids[s].add(e)
def clean(v):
 if v is None or (isinstance(v,float) and not math.isfinite(v)):return None
 return v.item() if hasattr(v,'item') else v
def mapping(e,s,source):
 e=str(e) if e is not None and not pd.isna(e) else ''; s=str(s) if s is not None and not pd.isna(s) else ''
 base=e.split('.')[0] if e.startswith('ENSG') else ''
 method='Ensembl_version_stripped' if base and base!=e else 'Ensembl_exact' if base else 'Unmatched'
 if not base and len(symids.get(s,set()))==1:base=next(iter(symids[s]));method='Symbol_exact'
 result={'original_gene_id':e or None,'original_gene_symbol':s or None,'harmonized_ensembl':base or None,'harmonized_symbol':ann.get(base,{}).get('symbol',s or None),'matching_method':method}
 if method=='Unmatched':unmatched.append({'source':source,**result,'reason':'No unambiguous mapping in existing GENCODE v26 annotation'})
 return result
def key(r):return r['harmonized_ensembl'] or 'SYMBOL:'+str(r['harmonized_symbol'])
def records(df):return [{k:clean(v) for k,v in r.items()} for r in df.to_dict('records')]
def readsmall(k,head):
 w=openpyxl.load_workbook(paths[k],read_only=True,data_only=True); s=w.active; vals=list(s.values); cols=list(vals[head-1]); last=max(i for i,x in enumerate(cols) if x is not None)+1; cols=cols[:last]
 counts=defaultdict(int); headers=[]
 for c in cols:counts[c]+=1;headers.append(str(c) if counts[c]==1 else str(c)+'_'+str(counts[c]))
 rows=[dict(zip(headers,r[:last])) for r in vals[head:] if any(x is not None for x in r[:last])]
 inventory.append({'file':names[k],'sheet':s.title,'physical_rows':s.max_row,'physical_columns':s.max_column,'header_row':head,'data_rows':len(rows),'columns':'; '.join(headers)})
 return rows
sets={}; backgrounds={}; own={}
for label in ['JTI','FUSION']:
 df=pd.read_csv(paths[label],sep='\t',low_memory=False); p='pvalue' if label=='JTI' else 'TWAS.P'; valid=pd.to_numeric(df[p],errors='coerce').between(0,1); sig=df[pd.to_numeric(df.FDR_ALL,errors='coerce')<.05].copy()
 backgrounds[label]=set(df.loc[valid,'ENSEMBL_ID'].str.split('.').str[0]); sets[label]=set(sig.ENSEMBL_ID.str.split('.').str[0]); own[label]={}
 for e,g in sig.groupby('ENSEMBL_ID'):
  r=g.iloc[0]; m=mapping(r['GENE_ID_ORIGINAL'] if label=='JTI' else r['ID'],r['GENE_SYMBOL'],label); own[label][e] = {**m,'min_FDR':float(g.FDR_ALL.min()),'significant_feature_count':len(g),'tissues':'; '.join(sorted(set(g['TISSUE' if label=='JTI' else 'PANEL'])))}
 inventory.append({'file':str(paths[label]),'sheet':'TSV','data_rows':len(df),'valid_P_rows':int(valid.sum()),'significant_rows':len(sig),'columns':'; '.join(df.columns)})
 print(label,len(df),len(sig),len(sets[label]),flush=True)
assert len(sets['JTI'])==64 and len(sets['FUSION'])==128
sets['Shared']=sets['JTI']&sets['FUSION'];sets['Union']=sets['JTI']|sets['FUSION'];assert len(sets['Shared'])==40
backgrounds['Shared']=backgrounds['JTI']&backgrounds['FUSION'];backgrounds['Union']=backgrounds['JTI']|backgrounds['FUSION']
focus='MAPT-AS1 KANSL1-AS1 KANSL1 MAPT LRRC37A LRRC37A2 CRHR1 ARHGAP27 ARL17A NSF PLEKHM1 SPPL2C FAM215B DND1P1'.split()
focusids={next(iter(symids[s])) for s in focus if len(symids[s])==1}; targets=sets['Union']|focusids
h=readsmall('HuangS1',3); prev=None
for r in h:
 for old,new in [('Beta','Replication_Beta'),('SE','Replication_SE'),('P','Replication_P'),('Beta_2','MetaAnalysis_Beta'),('SE_2','MetaAnalysis_SE'),('P_2','MetaAnalysis_P')]:r[new]=r.pop(old)
for i,r in enumerate(h,4):
 original=r['Gene']; prev=r['Gene'] or prev;r['Gene']=prev;r.update(mapping(None,prev,'HuangS1'));r.update({'source_excel_row':i,'source_gene_cell':original,'gene_fill_from_merged_cell':original is None})
assert len(h)==31
hg=[]
for k in sorted(set(map(key,h))):
 rr=[r for r in h if key(r)==k]; hg.append({**{c:rr[0][c] for c in ['original_gene_id','original_gene_symbol','harmonized_ensembl','harmonized_symbol','matching_method']},'feature_count':len(rr),'tissues':'; '.join(sorted(set(r['Tissue'] for r in rr)))})
fine=readsmall('HuangS3',2); finemap=defaultdict(list); fine_features=[]
for source_row,r in enumerate(fine,3):
 for entry in r['Gene'].split(','):
  gene=entry.strip().split(' ')[0]; m=mapping(None,gene,'HuangS3'); result=r['Significant (P < 0.05)\n(N = 7,805 trios)']; match=re.search(re.escape(entry.strip())+r' \(P=([^)]*)\)',result)
  finemap[key(m)].append({'P':float(match.group(1)) if match else None,'notes':result,'tissue':r['Tissue']})
  fine_features.append({**m,'Tissue':r['Tissue'],'CHR':r['CHR'],'source_feature':entry.strip(),'source_gene_list':r['Gene'],'source_excel_row':source_row,'Huang_finemapping_present':True,'Huang_finemapping_P':float(match.group(1)) if match else None,'Huang_finemapping_notes':result,'P_availability':'Reported significant conditional P' if match else 'Not supplied for this feature; not imputed'})
readsmall('HuangS2',3)
mat=readsmall('Matoba',1)
for r in mat:r.update(mapping(r['GENE'],r['hgnc_symbol'],'Matoba'))
assert len(mat)==567 and all(r['FDR']<.1 for r in mat)
print('Huang',len(hg),'Matoba',len(mat),flush=True)
wb=openpyxl.load_workbook(paths['Gandal'],read_only=True,data_only=True); gd=[]; secondary=[]; region=defaultdict(list); die=defaultdict(list); die_tested=set(); cytoband={}
for s in wb.worksheets:
 it=s.values; cols=next(it); n=0
 if s.title=='SupplementaryData3':
  n=sum(1 for _ in it);inventory.append({'file':names['Gandal'],'sheet':s.title,'physical_rows':s.max_row,'physical_columns':s.max_column,'data_rows':n,'columns':'Sheet; Description'});continue
 regcols=[c for c in cols if c.startswith('ASD_') and c.endswith('_FDR')]
 for rr in it:
  n+=1;r=dict(zip(cols,rr)); e=r['ensembl_gene_id'].split('.')[0]
  if s.title=='DEGene_Statistics':
   m=mapping(r['ensembl_gene_id'],r['external_gene_name'],'Gandal_DGE'); cytoband[e]=str(r['chromosome_name'])+str(r['band'])
   gd.append({**m,'gene_biotype':r['gene_biotype'],'chromosome':r['chromosome_name'],'band':r['band'],'logFC':r['WholeCortex_ASD_logFC'],'P':None,'FDR':r['WholeCortex_ASD_FDR'],'significant_FDR_lt_005':isinstance(r['WholeCortex_ASD_FDR'],(float,int)) and r['WholeCortex_ASD_FDR']<.05})
  elif e in targets:die_tested.add(e)
  if e not in targets:continue
  for c in regcols+(['WholeCortex_ASD_FDR'] if s.title=='DEIsoform_Statistics' else []):
   if isinstance(r[c],(int,float)) and r[c]<.05:
    typ='region_DGE' if s.title=='DEGene_Statistics' else 'isoform_DIE'; context=c[:-4]; (region if typ=='region_DGE' else die)[e].append(context)
    secondary.append({'Ensembl_ID':e,'Gene':ann.get(e,{}).get('symbol',r['external_gene_name']),'type':typ,'context':context,'transcript_id':r.get('ensembl_transcript_id'),'logFC':r[c.replace('_FDR','_logFC')],'FDR':r[c],'source_sheet':s.title,'source_excel_row':n+1,**mapping(r['ensembl_gene_id'],r['external_gene_name'],'Gandal_secondary')})
 inventory.append({'file':names['Gandal'],'sheet':s.title,'physical_rows':s.max_row,'physical_columns':s.max_column,'data_rows':n,'columns':'; '.join(cols)})
 print(s.title,n,flush=True)
gmap={key(r):r for r in gd}; assert len(gmap)==len(gd)
external={'Huang':set(map(key,hg)),'Matoba':set(map(key,mat)),'Gandal':{key(r) for r in gd if r['significant_FDR_lt_005']}}
hmap={key(r):r for r in hg};mmap={key(r):r for r in mat}
for s in list(sets['Union']):
 if cytoband.get(s)=='17q21.31':focusids.add(s)
# Include significant genes in the GRCh38 cytoband even if not represented in Gandal.
for e in sets['Union']:
 a=ann.get(e,{})
 if a.get('chr')=='chr17' and 42800000<=a.get('start',0)<46800000:focusids.add(e)
sfpath=ROOT/'outputs/restructure_20260922/Supplementary_tables/S13B_SFARI_overlap.tsv';fupath=ROOT/'outputs/restructure_20260922/Supplementary_tables/S14B_Fu2022_overlap.tsv'
paths['SFARI_existing']=sfpath;paths['Fu_existing']=fupath
sf=pd.read_csv(sfpath,sep='\t');fu=pd.read_csv(fupath,sep='\t'); sf1={r.Gene for _,r in sf.iterrows() if r['SFARI score']==1};sf12=set(sf.Gene); fu72=set(fu.loc[fu.ASD72==1,'ENSEMBL_ID']);fu185=set(fu.loc[fu.ASD185==1,'ENSEMBL_ID'])
we=set('NSF LRRC37A LRRC37A2 LRRC37A17P RNF123 RP11-156P1.3 MAPT-AS1'.split());ws=set('NT5C2 CRHR1 LRRC37A LRRC37A2 USMG5 KANSL1 USP4'.split()); walker=[]
for s in sorted(we|ws):walker.append({**mapping(None,s,'Walker_user_supplied'),'Gene':s,'Prenatal_expression_TWAS':s in we,'Prenatal_splicing_TWAS':s in ws,'Interpretation':'Developmental-brain / intracranial-volume context; not ASD replication','source':'User supplied predefined significant lists; no new threshold or enrichment'})
wemap={key(r) for r in walker if r['Prenatal_expression_TWAS']};wsmap={key(r) for r in walker if r['Prenatal_splicing_TWAS']}
matrix=[]
for e in sorted(sets['Union']|focusids,key=lambda e:ann.get(e,{}).get('symbol',e)):
 s=ann.get(e,{}).get('symbol',e);gg=gmap.get(e,{});hh=hmap.get(e,{});mm=mmap.get(e,{});ff=finemap.get(e,[])
 r={'Gene':s,'Ensembl_ID':e,'JTI_significant':e in sets['JTI'],'FUSION_significant':e in sets['FUSION'],'JTI_FUSION_shared':e in sets['Shared'],'JTI_min_FDR':own['JTI'].get(e,{}).get('min_FDR'),'FUSION_min_FDR':own['FUSION'].get(e,{}).get('min_FDR'),'SFARI_Category1':s in sf1,'SFARI_Category1_2':s in sf12,'Fu_ASD72':e in fu72,'Fu_ASD185':e in fu185,'Huang2021_TWAS':e in external['Huang'],'Huang2021_feature_count':hh.get('feature_count',0),'Huang2021_tissues':hh.get('tissues'),'Huang2021_finemapping_present':bool(ff),'Huang_finemapping_P':'; '.join(str(x['P']) for x in ff if x['P'] is not None) or None,'Huang_finemapping_notes':'; '.join(x['tissue']+': '+x['notes'] for x in ff) or None,'Matoba2020_HMAGMA':e in external['Matoba'],'Matoba2020_P':mm.get('P'),'Matoba2020_FDR':mm.get('FDR'),'Matoba2020_biotype':mm.get('gene_biotype'),'Gandal2022_whole_cortex_DGE':e in external['Gandal'],'Gandal2022_tested':e in gmap,'Gandal2022_logFC':gg.get('logFC'),'Gandal2022_P':None,'Gandal2022_FDR':gg.get('FDR'),'Gandal2022_region_specific_support':bool(region[e]),'Gandal2022_regions':'; '.join(sorted(set(region[e]))) or None,'Gandal2022_DIE_support':bool(die[e]),'Gandal2022_DIE_tested':e in die_tested,'Gandal2022_DIE_contexts':'; '.join(sorted(set(die[e]))) or None,'Walker2019_prenatal_expression':e in wemap,'Walker2019_prenatal_splicing':e in wsmap,'cytoband':cytoband.get(e),'focus_17q21_31':e in focusids}
 # Existing overlap tables cover only the significant union: do not assert negative SFARI/Fu membership outside it.
 if e not in sets['Union']:
  for c in ['SFARI_Category1','SFARI_Category1_2','Fu_ASD72','Fu_ASD185']:r[c]=None
 r['SFARI_source_support']=r['SFARI_Category1_2'];r['Fu_source_support']=r['Fu_ASD185'];r['Huang_source_support']=r['Huang2021_TWAS'];r['Matoba_source_support']=r['Matoba2020_HMAGMA'];r['Gandal_source_support']=r['Gandal2022_whole_cortex_DGE']
 r['N_ASD_external_sources_supported']=sum(r[c] is True for c in ['SFARI_source_support','Fu_source_support','Huang_source_support','Matoba_source_support','Gandal_source_support']);r['N_total_context_sources_supported']=r['N_ASD_external_sources_supported']+int(e in wemap or e in wsmap);r['source_count_complete']=e in sets['Union'];matrix.append(r)
def symbols(es):return '; '.join(sorted(ann.get(e,{}).get('symbol',e) for e in es))
overlap=[]
for label,ss in sets.items():
 for study,es in external.items():overlap.append({'TWAS_set':label,'external_study':study,'N_TWAS_set':len(ss),'N_external_set':len(es),'N_overlap':len(ss&es),'overlap_genes':symbols(ss&es),'overlap_Ensembl':'; '.join(sorted(ss&es)),'overlap_17q21_31':symbols(ss&es&focusids)})
enrich=[];gu={key(r) for r in gd if isinstance(r['FDR'],(int,float)) and 0<=r['FDR']<=1}
for label in sets:
 u=gu&backgrounds[label];a=sets[label]&u;b=external['Gandal']&u;x=len(a&b); cells=[x,len(a)-x,len(b)-x,len(u)-len(a|b)];od,p=fisher_exact([cells[:2],cells[2:]],alternative='greater')
 enrich.append({'TWAS_set':label,'study':'Gandal','external_testable_N':len(gu),'common_testable_N':len(u),'TWAS_significant_in_background':len(a),'external_significant_in_background':len(b),'a_overlap':cells[0],'b_TWAS_only':cells[1],'c_external_only':cells[2],'d_neither':cells[3],'odds_ratio':clean(float(od)),'Fisher_P':float(p)})
order=sorted(range(4),key=lambda i:enrich[i]['Fisher_P']);running=1
for rank in range(3,-1,-1):i=order[rank];running=min(running,enrich[i]['Fisher_P']*4/(rank+1));enrich[i]['BH_FDR_4_Gandal_tests']=running
notes=[
 ('Scope','Data整理、gene matching、descriptive overlap and source-specific enrichment only. No manuscript editing or new TWAS/SMR/coloc.'),
 ('Our sets','All-tissue BH-FDR < 0.05; JTI 64, FUSION 128, shared 40; union computed from identifiers.'),
 ('Matching','GENCODE v26 GRCh38 gene annotation. Strip Ensembl versions; unique exact symbol-to-Ensembl conversion only. Ambiguous/unresolved symbols retained as separate SYMBOL keys and listed in Unmatched_gene_records. No silent symbol-only overlap.'),
 ('Huang criterion','All 31 author-selected S1 meta-analysis significant features, including splicing; no filtering on discovery FDR column. Merged gene cells forward-filled with audit field.'),
 ('Huang independence','Family-based discovery: 7,805 AGP/SSC/SPARK trios. S1 meta-analysis includes 35,740 iPSYCH replication samples also included in Grove 2019. Therefore S1 is not fully independent from Grove; source counts are not independent replication counts.'),
 ('Matoba criterion','All 567 Table S7 genes, FDR < 0.1, all biotypes. ASD H-MAGMA regulatory gene mapping. Partially independent ASD cohort; orthogonal method.'),
 ('Gandal criterion','DEGene_Statistics: WholeCortex_ASD_FDR < 0.05 only; idiopathic ASD vs controls, gene-level whole cortex. All tested rows preserved. P values are not supplied.'),
 ('Gandal secondary','ASD region DGE and whole-cortex/region isoform DIE at source FDR < 0.05, descriptive only; selected union and focus genes. No dup15q columns. Any-region flag is not an additional cross-region corrected test.'),
 ('Enrichment Huang/Matoba','Formal enrichment was not performed because the full external testable gene universe could not be reconstructed reliably.'),
 ('Enrichment Gandal','One-sided Fisher, universe = valid whole-cortex ASD DGE genes intersected with corresponding valid-P TWAS-testable set; shared uses intersection, union uses union. BH across four Gandal comparisons. Four contingency cells retained. Gene correlations/LD are not modeled.'),
 ('Existing SFARI/Fu','Read existing S13/S14 results without redefining membership or correcting P values again. For focus genes outside our significant union, SFARI/Fu unassessed values remain blank.'),
 ('Counts','N_ASD counts SFARI Cat1-2, Fu ASD185, Huang S1, Matoba S7, Gandal primary, once per source. N_total adds Walker once. Secondary Gandal support alone is excluded from primary counts. Sources may share data.'),
 ('17q21.31','Requested focus list plus significant genes assigned 17q21.31 by Gandal annotation or located in GRCh38 chr17:42,800,000-46,800,000 by GENCODE. Coordinates used only for regional inventory, not gene matching.'),
 ('Walker','User-specified prenatal ICV significant expression/splicing list only; developmental context, no ASD enrichment. Unresolved legacy symbols remain unresolved.'),
 ('Interpretation','Published convergence / external molecular context / developmental-brain context. Absence means no support under this predefined source criterion, not evidence against a gene. Not tested is distinct from tested nonsignificant.'),
 ('Excluded','Huang S2 retained as provenance only. Sey 2020 methodological reference only. Pain/Grove FUSION not counted as independent evidence.'),
 ('Static analysis','Results computed reproducibly in Python from source data. Workbook is a fixed dated analysis snapshot; changing cells does not rerun set matching or Fisher tests.')]
readme=[{'Item':a,'Details':b} for a,b in notes]
studies=[
 ('Huang2021','Transcriptome-wide transmission disequilibrium analysis identifies novel risk genes for autism spectrum disorder',2021,'PLOS Genetics','Family-based transcriptomic association and case-control meta-analysis','7,805 trios + 35,740 iPSYCH samples','ASD','Author-selected S1 significant associations','Partly overlapping with Grove 2019 through iPSYCH','External family-based transcriptomic evidence','HuangS1','S1 Table','10.1371/journal.pgen.1009309'),
 ('Matoba2020','Common genetic risk variants identified in the SPARK cohort support DDHD2 as a candidate risk gene for autism',2020,'Translational Psychiatry','ASD H-MAGMA regulatory gene mapping','SPARK + prior GWAS meta-analysis; fetal-brain Hi-C','ASD','S7 all 567 genes, FDR < 0.1','Partially independent ASD cohort; orthogonal method','Regulatory gene mapping','Matoba','Table S7','10.1038/s41398-020-00953-9'),
 ('Gandal2022','Broad transcriptomic dysregulation occurs across the cerebral cortex in ASD',2022,'Nature','Direct postmortem cortex RNA-seq DGE','ASD and control postmortem cerebral cortex','Idiopathic ASD vs control','Whole-cortex gene-level FDR < 0.05','Independent molecular samples; no GWAS gene mapping','External molecular context','Gandal','DEGene_Statistics','10.1038/s41586-022-05377-7'),
 ('Walker2019','Genetic control of expression and splicing in developing human brain informs disease mechanisms',2019,'Cell','Prenatal brain TWAS','Developing brain expression/splicing models and ICV GWAS','Intracranial volume','User-specified published significant gene lists','Not ASD replication','Developmental context only',None,'User supplied list','10.1016/j.cell.2019.09.021')]
provenance=[]
for st,title,year,journal,dt,design,phen,criterion,ind,cat,k,sh,doi in studies:
 row=dict(zip(['Study','Article','Year','Journal','Data_type','Sample_design','Main_phenotype','Gene_selection_criterion','Independence_from_Grove2019','Evidence_category','Input_file','Sheet_used','Article_URL'],[st,title,year,journal,dt,design,phen,criterion,ind,cat,str(paths[k]) if k else 'User task attachment',sh,'https://doi.org/'+doi]));provenance.append(row)
 for field in ['Article','Year','Journal','Data_type','Sample_design','Main_phenotype','Gene_selection_criterion','Independence_from_Grove2019','Evidence_category','Input_file','Sheet_used']:readme.append({'Item':st+' / '+field,'Details':str(row[field])})
for k,p in paths.items():provenance.append({'Study':k,'Input_file':str(p),'SHA256':hashlib.sha256(p.read_bytes()).hexdigest(),'File_bytes':p.stat().st_size})
sheets['README']=readme
sheets['Our_JTI_64']=list(own['JTI'].values());sheets['Our_FUSION_128']=list(own['FUSION'].values());sheets['Our_Shared_40']=[r for r in matrix if r['Ensembl_ID'] in sets['Shared']];sheets['Our_Union_152']=[r for r in matrix if r['Ensembl_ID'] in sets['Union']]
sheets['Huang2021_features']=h;sheets['Huang2021_genes']=hg;sheets['Huang2021_finemapping']=fine_features;sheets['Matoba2020_genes']=mat;sheets['Gandal2022_whole_cortex_DGE']=gd;sheets['Gandal2022_secondary_support']=secondary;sheets['Walker2019_context']=walker;sheets['Overlap_summary']=overlap;sheets['External_Evidence_Matrix']=[r for r in matrix if r['Ensembl_ID'] in sets['Union']];sheets['ExternalEvidence_17q21.31']=[r for r in matrix if r['Ensembl_ID'] in focusids]
sheets['Unmatched_gene_records']=list({json.dumps(r,sort_keys=True):r for r in unmatched}.values()) or [{'Status':'No unmatched records'}];sheets['Provenance']=provenance;sheets['Workbook_inventory']=inventory;sheets['Gandal_enrichment']=enrich
for name in ['S13A_SFARI_enrichment','S14A_Fu2022_enrichment']:
 p=ROOT/'outputs/restructure_20260922/Supplementary_tables'/str(name+'.tsv');sheets[name]=records(pd.read_csv(p,sep='\t'));provenance.append({'Study':'Existing '+name,'Input_file':str(p),'SHA256':hashlib.sha256(p.read_bytes()).hexdigest()})
def normalize(obj):
 if isinstance(obj,dict):return {k:normalize(v) for k,v in obj.items()}
 if isinstance(obj,list):return [normalize(v) for v in obj]
 return clean(obj)
(WORK/'tables.json').write_text(json.dumps(normalize(sheets),ensure_ascii=False,allow_nan=False),encoding='utf8')
report=['ASD External Evidence Benchmark — 2026-09-26','', 'Fixed analysis snapshot. No manuscript edits.','', '1. WORKBOOK INVENTORY',json.dumps(inventory,ensure_ascii=False,indent=2),'','2. OUR SETS',str({k:len(v) for k,v in sets.items()}),'','3. EXTERNAL SETS',str({k:len(v) for k,v in external.items()}),'Huang feature count: 31. Gandal tested universe: '+str(len(gu)),'','4–6. OVERLAPS AND 17q21.31',pd.DataFrame(overlap).to_string(index=False),'','7. MAPT-AS1',json.dumps([r for r in matrix if r['Gene']=='MAPT-AS1'],ensure_ascii=False,indent=2),'','8. FOCUS GENES WITH NO PRIMARY ASD SOURCE SUPPORT',symbols({r['Ensembl_ID'] for r in matrix if r['Ensembl_ID'] in focusids and r['N_ASD_external_sources_supported']==0}), 'No primary support is not equivalent to no descriptive regional/isoform or developmental context.','Focus genes with no primary, secondary or developmental support: '+symbols({r['Ensembl_ID'] for r in matrix if r['Ensembl_ID'] in focusids and r['N_total_context_sources_supported']==0 and not r['Gandal2022_region_specific_support'] and not r['Gandal2022_DIE_support']}),'','9. UNRESOLVED MATCHING',json.dumps(sheets['Unmatched_gene_records'],ensure_ascii=False,indent=2),'','10. ENRICHMENT',pd.DataFrame(enrich).to_string(index=False),'','METHODS / LIMITATIONS']
report += [a+': '+b for a,b in notes]
report += ['','REFERENCES']+[str(r.get('Article',''))+' '+str(r.get('Article_URL','')) for r in provenance[:4]]
(OUT/'ASD_External_Evidence_Benchmark_Report_2026-09-26.txt').write_text('\n'.join(report),encoding='utf-8-sig')
print(pd.DataFrame(overlap)[['TWAS_set','external_study','N_external_set','N_overlap']].to_string(index=False),flush=True)
print('UNMATCHED',sheets['Unmatched_gene_records'],flush=True)
print('MAPT-AS1', [r for r in matrix if r['Gene']=='MAPT-AS1'],flush=True)
